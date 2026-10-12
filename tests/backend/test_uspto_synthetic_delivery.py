"""P2-007 controlled synthetic USPTO download API contract and security checks."""
from __future__ import annotations

from io import BytesIO
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
import pytest

from backend.app.demo_staging import create_demo_staging_app
from backend.app.main import create_app
from backend.app.services.uspto_offline_renderer import BANNER, CLAIMS_INCOMPLETE, FILES
from backend.app.services.uspto_synthetic_delivery import (
    FIXTURE, MIME_TYPES, SAMPLE_ID, SyntheticUSPTOExports,
)


@pytest.fixture
def client(tmp_path):
    with TestClient(create_demo_staging_app(temporary_root=tmp_path)) as session:
        yield session


def _body_text(content: bytes, name: str) -> str:
    if name.endswith(".pdf"):
        return " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    document = Document(BytesIO(content))
    return " ".join([*(p.text for p in document.paragraphs),
                     *(p.text for section in document.sections for p in section.header.paragraphs)])


def test_capability_manifest_only_after_validated_generation(client):
    response = client.get("/api/demo/uspto/exports")
    assert response.status_code == 200
    assert response.headers["cache-control"] == "no-store"
    manifest = response.json()
    assert manifest["sample_id"] == SAMPLE_ID
    assert manifest["source"] == "controlled_synthetic_p2_005"
    assert (manifest["draft_only"], manifest["filing_ready"], manifest["claims_complete"]) == (True, False, False)
    assert [(item["filename"], item["media_type"]) for item in manifest["files"]] == list(MIME_TYPES.items())


@pytest.mark.parametrize("name", FILES)
def test_real_four_documents_signatures_content_and_headers(client, name):
    response = client.get(f"/api/demo/uspto/exports/{SAMPLE_ID}/{name}")
    assert response.status_code == 200
    assert response.headers["content-type"] == MIME_TYPES[name]
    assert response.headers["content-disposition"] == f'attachment; filename="{name}"'
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert len(response.content) > 1000
    assert response.content.startswith(b"%PDF-" if name.endswith(".pdf") else b"PK\x03\x04")
    content = re.sub(r"\s+", " ", _body_text(response.content, name))
    assert BANNER in content and CLAIMS_INCOMPLETE in content
    assert "What is claimed is:" not in content
    if name.endswith(".docx"):
        with ZipFile(BytesIO(response.content)) as archive:
            assert 'word/document.xml' in archive.namelist()
            assert not any('vbaProject' in key for key in archive.namelist())


@pytest.mark.parametrize("sample,name", [
    ("unknown", FILES[0]),
    (SAMPLE_ID, "random.pdf"),
    (SAMPLE_ID, ".env"),
    ("..", FILES[0]),
    (SAMPLE_ID, "B2_REFERENCE_PREVIEW.pdf.bak"),
])
def test_fixed_allowlist_refuses_other_samples_and_files(client, sample, name):
    response = client.get(f"/api/demo/uspto/exports/{sample}/{name}", follow_redirects=False)
    assert response.status_code in (404, 400)
    assert b"Traceback" not in response.content


def test_encoded_path_traversal_and_no_untrusted_body(client):
    for path in (
        f"/api/demo/uspto/exports/{SAMPLE_ID}/%2e%2e%2fAGENTS.md",
        f"/api/demo/uspto/exports/{SAMPLE_ID}/%2e%2e%5cAGENTS.md",
        f"/api/demo/uspto/exports/{SAMPLE_ID}/%252e%252e%252fAGENTS.md",
    ):
        response = client.get(path)
        assert response.status_code in (400, 404)
    response = client.post(f"/api/demo/uspto/exports/{SAMPLE_ID}/{FILES[0]}", json={"package": {"claims": ["fake"]}})
    assert response.status_code == 405


def test_tampered_fixture_fails_closed_without_internal_information(client, monkeypatch, tmp_path):
    import backend.app.services.uspto_synthetic_delivery as delivery
    tampered = tmp_path / "representative-system.json"
    tampered.write_bytes(FIXTURE.read_bytes() + b" ")
    monkeypatch.setattr(delivery, "FIXTURE", tampered)
    response = client.get("/api/demo/uspto/exports")
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "synthetic_export_unavailable"
    assert "fixture" not in response.text.lower()
    assert str(tmp_path) not in response.text
    assert client.get(f"/api/demo/uspto/exports/{SAMPLE_ID}/{FILES[0]}").status_code == 503


def test_no_uspto_route_on_normal_backend():
    normal = TestClient(create_app())
    assert normal.get("/api/demo/uspto/exports").status_code == 404
    assert normal.get(f"/api/demo/uspto/exports/{SAMPLE_ID}/{FILES[0]}").status_code == 404


def test_renderer_failure_returns_generic_unavailable_without_content(client, monkeypatch):
    import backend.app.services.uspto_synthetic_delivery as delivery

    def broken_renderer(*args, **kwargs):
        raise RuntimeError("sensitive renderer detail")

    monkeypatch.setattr(delivery, "render_uspto_document_package", broken_renderer)
    for endpoint in (
        "/api/demo/uspto/exports",
        f"/api/demo/uspto/exports/{SAMPLE_ID}/{FILES[0]}",
    ):
        response = client.get(endpoint)
        assert response.status_code == 503
        assert response.json()["error"]["code"] == "synthetic_export_unavailable"
        assert "sensitive" not in response.text


def test_cache_generates_only_once(monkeypatch):
    delivery = SyntheticUSPTOExports()
    count = 0
    original = delivery._generate
    def counted():
        nonlocal count
        count += 1
        return original()
    monkeypatch.setattr(delivery, "_generate", counted)
    assert delivery.get_files() is delivery.get_files()
    assert count == 1
