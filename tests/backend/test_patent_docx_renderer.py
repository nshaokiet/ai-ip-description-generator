from __future__ import annotations

import copy
import hashlib
import json
import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from jsonschema import ValidationError

from backend.app.services.patent_docx_renderer import (
    TEMPLATE_PATH,
    PatentDocxRenderer,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
FIXTURE_ROOT = REPOSITORY_ROOT / "tests/fixtures/patent-docx/v1.0.0"
CONFLICT_CASE = (
    REPOSITORY_ROOT
    / "tests/fixtures/patent-draft/v1.0.0/cases/conflicting-evidence.json"
)
SECTION_HEADINGS = [
    "Cover and Draft Notice",
    "Title",
    "Abstract",
    "Technical Field",
    "Background",
    "Summary of Technical Disclosure",
    "Brief Description of Drawings",
    "Detailed Description",
    "Experimental Examples and Results",
    "Missing Information and Author Questions",
    "Evidence and Reference Appendix",
]


def _load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _all_text(path: Path) -> str:
    document = Document(path)
    values = [paragraph.text for paragraph in document.paragraphs]
    for section in document.sections:
        values.extend(paragraph.text for paragraph in section.header.paragraphs)
        values.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(values)


def _package_text(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        return "\n".join(
            package.read(name).decode("utf-8", errors="replace")
            for name in package.namelist()
            if name.endswith((".xml", ".rels"))
        )


class PatentDocxRendererProductionTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temporary_directory = tempfile.TemporaryDirectory()
        self.output_root = Path(self.temporary_directory.name) / "approved-output"
        self.renderer = PatentDocxRenderer(output_root=self.output_root)
        self.representative = _load_json(FIXTURE_ROOT / "representative-context.json")

    def tearDown(self) -> None:
        self.temporary_directory.cleanup()

    def test_valid_canonical_draft_renders_with_canonical_template_offline(self) -> None:
        output = self.output_root / "representative.docx"
        template_hash = hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
        with (
            patch("socket.create_connection") as connect,
            patch("urllib.request.urlopen") as urlopen,
        ):
            rendered = self.renderer.render(self.representative, output)

        connect.assert_not_called()
        urlopen.assert_not_called()
        self.assertEqual(output.resolve(), rendered)
        self.assertTrue(zipfile.is_zipfile(output))
        with zipfile.ZipFile(output) as package:
            self.assertIsNone(package.testzip())
        self.assertEqual(template_hash, hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest())
        headings = [
            paragraph.text
            for paragraph in Document(output).paragraphs
            if paragraph.style.name == "Heading 1"
        ]
        self.assertEqual(SECTION_HEADINGS, headings)

    def test_invalid_schema_payload_is_rejected_without_output(self) -> None:
        invalid = copy.deepcopy(self.representative)
        invalid["claims"] = []
        output = self.output_root / "invalid.docx"
        with self.assertRaises(ValidationError):
            self.renderer.render(invalid, output)
        self.assertFalse(output.exists())

    def test_automatic_escaping_preserves_literal_template_like_data(self) -> None:
        payload = _load_json(FIXTURE_ROOT / "special-characters-context.json")
        output = self.output_root / "special.docx"
        self.renderer.render(payload, output)
        text = _all_text(output)
        self.assertIn("<sensor & controller>", text)
        self.assertIn("{{ do_not_execute }}", text)
        self.assertIn("{% raw_instruction %}", text)
        for unresolved in ("{{ paragraph.", "{{ evidence.", "{%p"):
            self.assertNotIn(unresolved, _package_text(output))

    def test_missing_marker_and_author_questions_are_preserved(self) -> None:
        payload = _load_json(FIXTURE_ROOT / "missing-information-context.json")
        output = self.output_root / "missing.docx"
        self.renderer.render(payload, output)
        text = _all_text(output)
        self.assertGreaterEqual(text.count("[AUTHOR CONFIRMATION REQUIRED]"), 10)
        self.assertIn("Which existing title candidate should be selected?", text)

    def test_conflict_inference_and_evidence_appendix_are_preserved(self) -> None:
        conflict = _load_json(CONFLICT_CASE)["expected_output"]
        conflict_output = self.output_root / "conflict.docx"
        representative_output = self.output_root / "grounding.docx"
        self.renderer.render(conflict, conflict_output)
        self.renderer.render(self.representative, representative_output)
        conflict_text = _all_text(conflict_output)
        representative_text = _all_text(representative_output)
        self.assertIn("UNRESOLVED", conflict_text)
        self.assertIn("ev-conflict-a, ev-conflict-b", conflict_text)
        self.assertIn("Grounding: INFERENCE | Evidence: ev-core", representative_text)
        self.assertIn("Evidence and Reference Appendix", representative_text)

    def test_output_destination_is_confined_to_configured_root(self) -> None:
        outside = Path(self.temporary_directory.name) / "outside.docx"
        traversal = self.output_root / ".." / "traversal.docx"
        for candidate in (outside, traversal):
            with self.subTest(candidate=candidate):
                with self.assertRaisesRegex(ValueError, "configured output root"):
                    self.renderer.render(self.representative, candidate)
                self.assertFalse(candidate.resolve().exists())

    def test_canonical_template_and_non_docx_destinations_are_rejected(self) -> None:
        template_renderer = PatentDocxRenderer(output_root=TEMPLATE_PATH.parent)
        with self.assertRaisesRegex(ValueError, "canonical template"):
            template_renderer.render(self.representative, TEMPLATE_PATH)
        with self.assertRaisesRegex(ValueError, r"\.docx extension"):
            self.renderer.render(self.representative, self.output_root / "draft.json")

    def test_claims_and_legal_conclusions_remain_absent(self) -> None:
        output = self.output_root / "claims-excluded.docx"
        self.renderer.render(self.representative, output)
        package_text = _package_text(output)
        document_text = _all_text(output).lower()
        self.assertIsNone(re.search(r"\b(?:draft\s+)?claims?\b", package_text, re.I))
        for prohibited in (
            "is patentable",
            "is novel",
            "has an inventive step",
            "ready for filing",
        ):
            self.assertNotIn(prohibited, document_text)


if __name__ == "__main__":
    unittest.main()
