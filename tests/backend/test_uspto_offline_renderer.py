"""P2-006 synthetic-only, seven-input public-validation and four-artifact tests."""
from __future__ import annotations

from copy import deepcopy
import re
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from lxml import etree
from pypdf import PdfReader

from backend.app.services.uspto_document_package_validation import USPTODocumentPackageValidationError
from backend.app.services.uspto_offline_renderer import (
    BANNER, CLAIMS_INCOMPLETE, FILES, Part, RenderModel, Section, _docx, _pdf,
    render_uspto_document_package,
)
from test_uspto_document_package_v1 import inputs, run, standalone_positive_projection


def _render(tmp_path: Path, original: dict | None = None):
    original = inputs() if original is None else original
    package = run(original)
    files = render_uspto_document_package(package, **original, output_dir=tmp_path)
    return original, package, files


def _docx_text(path: Path) -> str:
    doc = Document(path)
    pieces = [p.text for p in doc.paragraphs]
    pieces.extend(p.text for section in doc.sections for p in section.header.paragraphs)
    for table in doc.tables:
        pieces.extend(cell.text for row in table.rows for cell in row.cells)
    return "\n".join(pieces)


def _pdf_text(path: Path) -> str:
    # Visual line wrapping is not a semantic difference.
    return re.sub(r"\s+", " ", "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages))


def test_four_mandatory_formats_order_status_and_filing_boundary(tmp_path):
    bundle, package, files = _render(tmp_path)
    assert tuple(files) == FILES
    assert all(p.is_file() and p.stat().st_size > 1000 for p in files.values())
    assert len({str(p.parent) for p in files.values()}) == 1
    assert all(p.parent.name == package["package_id"] + "-p2-006" for p in files.values())
    leaked = ["review_package", "source_paths", "evidence_ids", "projected_claim_support",
              "original_p2_003", "original_p2_004", "MI-P2-005-0001"]
    for name, path in files.items():
        text = _pdf_text(path) if path.suffix == ".pdf" else _docx_text(path)
        assert BANNER in text
        assert CLAIMS_INCOMPLETE in text
        assert "What is claimed is:" not in text
        assert "1. A " not in text
        assert all(term not in text for term in leaked)
        assert text.index("TITLE OF THE INVENTION") < text.index("BACKGROUND OF THE INVENTION")
        if path.suffix == ".pdf":
            assert len(PdfReader(str(path)).pages) >= 2
            for page in PdfReader(str(path)).pages:
                assert BANNER in page.extract_text()
        else:
            assert Document(path).paragraphs
            assert not any(item.lower().endswith("vbaproject.bin") for item in ZipFile(path).namelist())
    assert package["filing_view"]["claims"]["status"] == "HUMAN_CONFIRMATION_REQUIRED"
    assert bundle["specification"]["review_boundary"]["claims_runtime_implemented"] is False


def test_pdf_embedded_font_geometry_and_application_new_sheets(tmp_path):
    _, _, files = _render(tmp_path)
    for filename in (FILES[0], FILES[2]):
        reader = PdfReader(str(files[filename]))
        assert all(float(p.mediabox.width) == 612 and float(p.mediabox.height) == 792 for p in reader.pages)
        for page in reader.pages:
            fonts = page["/Resources"]["/Font"].get_object().values()
            assert any("/FontFile2" in font.get_object().get("/FontDescriptor", {}).get_object()
                       for font in fonts if font.get_object().get("/FontDescriptor"))
    application = PdfReader(str(files[FILES[2]]))
    claims_pages = [i for i, p in enumerate(application.pages)
                    if CLAIMS_INCOMPLETE in re.sub(r"\s+", " ", p.extract_text())]
    abstract_pages = [i for i, p in enumerate(application.pages) if "ABSTRACT OF THE DISCLOSURE" in p.extract_text()]
    assert claims_pages and abstract_pages
    assert claims_pages[0] < abstract_pages[0]


def test_docx_editable_columns_line_numbers_and_single_column_application(tmp_path):
    _, _, files = _render(tmp_path)
    namespace = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with ZipFile(files[FILES[1]]) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        assert root.xpath("//w:sectPr/w:cols[@w:num='2']", namespaces=namespace)
        assert root.xpath("//w:sectPr/w:lnNumType[@w:countBy='5']", namespaces=namespace)
        assert root.xpath("//w:t", namespaces=namespace)
    with ZipFile(files[FILES[3]]) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        assert not root.xpath("//w:sectPr/w:cols[@w:num='2']", namespaces=namespace)
        assert len(root.xpath("//w:br[@w:type='page']", namespaces=namespace)) >= 2


@pytest.mark.parametrize("mutation", ["package", "source", "claims", "config"])
def test_mutation_fails_before_any_output(tmp_path, mutation):
    originals = inputs()
    package = run(originals)
    if mutation == "package":
        package["package_id"] = "p2-005-" + "0" * 24
    elif mutation == "source":
        originals["structured_analysis"]["unapproved_key"] = "mutation"
    elif mutation == "claims":
        package["filing_view"]["claims"]["claims"] = [{"claim_text": "fabricated"}]
    else:
        originals["p2_004_generation_config"]["max_dependent_claims_per_family"] = 99
    with pytest.raises(USPTODocumentPackageValidationError):
        render_uspto_document_package(package, **originals, output_dir=tmp_path)
    assert list(tmp_path.iterdir()) == []


def test_absent_sidecar_and_destination_collision_fail_closed(tmp_path):
    original = inputs()
    package = run(original)
    arguments = deepcopy(original)
    arguments.pop("source_envelope")
    with pytest.raises(TypeError):
        render_uspto_document_package(package, **arguments, output_dir=tmp_path)
    assert list(tmp_path.iterdir()) == []
    _, _, files = _render(tmp_path, original)
    with pytest.raises(FileExistsError):
        render_uspto_document_package(package, **original, output_dir=tmp_path)
    assert all(path.exists() for path in files.values())


def test_no_unvalidated_component_positive_claims_output(tmp_path):
    projected, claims = standalone_positive_projection()
    assert claims["claims"]
    with pytest.raises(USPTODocumentPackageValidationError):
        render_uspto_document_package(projected, **inputs(), output_dir=tmp_path)
    assert list(tmp_path.iterdir()) == []


def test_no_external_ole_macro_or_links(tmp_path):
    _, _, files = _render(tmp_path)
    for name in (FILES[1], FILES[3]):
        with ZipFile(files[name]) as archive:
            members = archive.namelist()
            assert not any("embeddings/" in item or "vbaProject" in item for item in members)
            rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
            assert 'TargetMode="External"' not in rels


def test_all_four_formats_preserve_verified_title_abstract_and_numeric_facts(tmp_path):
    _, package, files = _render(tmp_path)
    filing = package["filing_view"]
    title = filing["sections"][0]["paragraphs"][0]["text"]
    abstract = filing["abstract"]["text"]
    assert filing["abstract"]["word_count"] <= 150
    for path in files.values():
        document_text = (_pdf_text(path) if path.suffix == ".pdf" else
                         re.sub(r"\s+", " ", _docx_text(path)))
        assert title in document_text
        assert abstract in document_text
        assert "25 degC" in document_text
        assert "12 percent" in document_text


def test_internal_layout_unicode_markup_and_multipage(tmp_path):
    # Synthetic internal layout probe; bypasses validator only to exercise
    # layout boundaries, never the public package rendering entrypoint.
    fragment = 'Tiếng Việt: nhiệt độ 25°C & hiệu suất <12%> — "đo lường".'
    model = RenderModel(
        title='Synthetic multilingual layout QA',
        abstract=fragment,
        sections=(
            Section('TITLE_OF_THE_INVENTION', 'TITLE OF THE INVENTION',
                    (Part('prose', fragment),)),
            Section('DETAILED_DESCRIPTION_OF_THE_INVENTION',
                    'DETAILED DESCRIPTION OF THE INVENTION',
                    tuple(Part('prose', f'Synthetic paragraph {i}: {fragment}')
                          for i in range(160))),
        ),
        draft_id='SYNTHETIC-QA',
    )
    for label, b2 in [('b2', True), ('application', False)]:
        pdf = tmp_path / f'{label}.pdf'
        docx = tmp_path / f'{label}.docx'
        _pdf(model, pdf, b2=b2)
        _docx(model, docx, b2=b2)
        assert len(PdfReader(str(pdf)).pages) >= 3
        assert len(_docx_text(docx)) > 1000
        assert fragment in _docx_text(docx)
        extracted = _pdf_text(pdf)
        assert fragment in extracted
        with ZipFile(docx) as archive:
            assert archive.testzip() is None
