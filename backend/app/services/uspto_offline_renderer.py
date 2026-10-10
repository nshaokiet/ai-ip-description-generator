"""P2-006: offline, synthetic-only four-format render of a validated P2-005 package.

This service is intentionally not exposed as a public API. The released P2-005
validator and its seven original inputs are the sole entry trust boundary.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from tempfile import mkdtemp
from typing import Any, Mapping
from xml.sax.saxutils import escape
import os
import shutil

from backend.app.services.uspto_document_package_validation import validate_uspto_document_package


BANNER = "PATENT SPECIFICATION DRAFT — NOT ISSUED"
CLAIMS_INCOMPLETE = "[CLAIMS INCOMPLETE — HUMAN CONFIRMATION REQUIRED; NO CLAIM TEXT SUPPLIED]"
ABSTRACT_INCOMPLETE = "[ABSTRACT INCOMPLETE — HUMAN CONFIRMATION REQUIRED]"
FILES = (
    "B2_REFERENCE_PREVIEW.pdf", "B2_REFERENCE_PREVIEW.docx",
    "USPTO_APPLICATION_DRAFT.pdf", "USPTO_APPLICATION_DRAFT.docx",
)
FONT_DIR = Path(__file__).resolve().parent / "assets/fonts"
MAX_FILING_CHARS = 1_000_000


@dataclass(frozen=True)
class Part:
    kind: str  # heading, prose, status, claim
    text: str


@dataclass(frozen=True)
class Section:
    section_id: str
    heading: str
    parts: tuple[Part, ...]


@dataclass(frozen=True)
class RenderModel:
    title: str
    abstract: str
    sections: tuple[Section, ...]
    draft_id: str


def _status_marker(status: str) -> str:
    if status == "NOT_APPLICABLE":
        return "[SECTION NOT APPLICABLE — VALIDATED BASIS]"
    return "[SECTION INCOMPLETE — HUMAN CONFIRMATION REQUIRED]"


def _parts(value: Mapping[str, Any]) -> tuple[Part, ...]:
    status = value["status"]
    if status != "PRESENT":
        return (Part("status", _status_marker(status)),)
    parts = [Part("prose", p["text"]) for p in value["paragraphs"]]
    # Only the controlled, released section hierarchy is traversed.
    for child in value.get("subsections", ()):
        parts.append(Part("heading", child["heading"]))
        parts.extend(_parts(child))
    return tuple(parts)


def _build_model(filing: Mapping[str, Any]) -> RenderModel:
    lookup = {section["section_id"]: section for section in filing["sections"]}
    title_section = lookup["TITLE_OF_THE_INVENTION"]
    title = " ".join(part["text"] for part in title_section["paragraphs"])
    if title_section["status"] != "PRESENT":
        title = _status_marker(title_section["status"])
    abstract_section = filing["abstract"]
    abstract = (abstract_section["text"] if abstract_section["status"] == "PRESENT"
                else ABSTRACT_INCOMPLETE)
    sections = []
    for section_id in filing["section_order"]:
        if section_id == "CLAIMS":
            claims_section = filing["claims"]
            if claims_section["status"] == "PRESENT":
                parts = [Part("prose", claims_section["preamble"])]
                parts.extend(Part("claim", f'{claim["claim_number"]}. {claim["claim_text"]}')
                             for claim in claims_section["claims"])
            else:
                parts = [Part("status", CLAIMS_INCOMPLETE)]
            sections.append(Section(section_id, claims_section["heading"], tuple(parts)))
        elif section_id == "ABSTRACT_OF_THE_DISCLOSURE":
            sections.append(Section(section_id, "ABSTRACT OF THE DISCLOSURE",
                                    (Part("prose" if abstract_section["status"] == "PRESENT"
                                          else "status", abstract),)))
        else:
            source = lookup[section_id]
            sections.append(Section(section_id, source["heading"], _parts(source)))
    return RenderModel(title, abstract, tuple(sections), filing["document_metadata"]["draft_id"])


def _assert_safe_parent(parent: Path) -> Path:
    # Caller chooses an existing private/local directory. Refuse symlink ancestors.
    if not parent.is_absolute() or not parent.is_dir():
        raise ValueError("A pre-existing absolute output directory is required")
    if any(path.is_symlink() for path in (parent, *parent.parents)):
        raise ValueError("Symlinked output directories are not accepted")
    return parent.resolve(strict=True)


def render_uspto_document_package(
    package: Mapping[str, Any], *, source_envelope: Any, structured_analysis: Any,
    reconstruction: Any, specification: Any, p2_003_generation_config: Any,
    supported_claims: Any, p2_004_generation_config: Any, output_dir: Path | str,
) -> dict[str, Path]:
    """Validate first, then atomically publish a four-file synthetic QA directory.

    Returns filenames mapped to absolute paths inside <output_dir>/<package_id>-p2-006.
    No package-only or caller-attested validation route is supported.
    """
    validate_uspto_document_package(
        package, source_envelope=source_envelope, structured_analysis=structured_analysis,
        reconstruction=reconstruction, specification=specification,
        p2_003_generation_config=p2_003_generation_config,
        supported_claims=supported_claims, p2_004_generation_config=p2_004_generation_config,
    )
    # All content consumed by the renderers is copied out of the filing-only view.
    from copy import deepcopy
    model = _build_model(deepcopy(package["filing_view"]))
    if sum(len(part.text) for section in model.sections for part in section.parts) > MAX_FILING_CHARS:
        raise ValueError("Validated filing content exceeds offline renderer limit")
    parent = _assert_safe_parent(Path(output_dir))
    destination = parent / (package["package_id"] + "-p2-006")
    if destination.exists():
        raise FileExistsError("Renderer output directory already exists")
    staging = Path(mkdtemp(prefix=".p2-006-", dir=parent))
    try:
        _pdf(model, staging / FILES[0], b2=True)
        _docx(model, staging / FILES[1], b2=True)
        _pdf(model, staging / FILES[2], b2=False)
        _docx(model, staging / FILES[3], b2=False)
        if set(item.name for item in staging.iterdir()) != set(FILES):
            raise RuntimeError("Mandatory four-output render failed")
        os.rename(staging, destination)
    finally:
        if staging.exists():
            shutil.rmtree(staging)
    return {name: destination / name for name in FILES}


def _pdf(model: RenderModel, path: Path, *, b2: bool) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (BaseDocTemplate, Flowable, Frame, KeepTogether,
                                   NextPageTemplate, PageBreak, PageTemplate, Paragraph,
                                   Spacer, Table, TableStyle)

    if not FONT_DIR.joinpath("OFL.txt").is_file():
        raise RuntimeError("Licensed font manifest unavailable")
    for label, name in (("CrimsonText", "CrimsonText-Regular.ttf"),
                        ("CrimsonText-Bold", "CrimsonText-Bold.ttf")):
        if label not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(label, str(FONT_DIR / name)))
    pdfmetrics.registerFontFamily("CrimsonText", normal="CrimsonText", bold="CrimsonText-Bold")

    size = 9.4 if b2 else 12.0
    leading = 11.6 if b2 else 18.0
    heading = ParagraphStyle("Heading", fontName="CrimsonText-Bold", fontSize=11 if b2 else 12,
                             leading=14 if b2 else 18, alignment=TA_CENTER, spaceBefore=12,
                             spaceAfter=7, keepWithNext=1)
    subheading = ParagraphStyle("Subheading", parent=heading, fontSize=9.6 if b2 else 12,
                                leading=12 if b2 else 18, spaceBefore=8)
    body = ParagraphStyle("Body", fontName="CrimsonText", fontSize=size, leading=leading,
                          alignment=TA_JUSTIFY, spaceAfter=5 if b2 else 9,
                          firstLineIndent=9 if b2 else 0, splitLongWords=1)
    status = ParagraphStyle("Status", parent=body, fontName="CrimsonText-Bold",
                            alignment=TA_LEFT, firstLineIndent=0,
                            spaceBefore=4, spaceAfter=8)
    claim = ParagraphStyle("Claim", parent=body, firstLineIndent=0, leftIndent=8,
                           spaceBefore=7, allowOrphans=0, allowWidows=0)
    cover_label = ParagraphStyle("CoverLabel", parent=heading, alignment=TA_CENTER, fontSize=11)
    cover_text = ParagraphStyle("CoverText", parent=body, fontSize=11, leading=14,
                                firstLineIndent=0)

    class Template(BaseDocTemplate):
        def afterFlowable(self, flowable):
            # Record the lowest rendered text position per column. Draw line
            # references only for populated portions of each column/page.
            if not b2 or self.page <= 1 or not isinstance(flowable, Paragraph):
                return
            if self.frame.id not in ("left", "right"):
                return
            if not hasattr(self, "_column_extents"):
                self._column_extents = {}
            self._column_extents[self.frame.id] = min(
                self._column_extents.get(self.frame.id, 726), self.frame._y
            )

    doc = Template(str(path), pagesize=letter, leftMargin=0, rightMargin=0,
                   topMargin=0, bottomMargin=0, title=BANNER,
                   author="Synthetic offline draft renderer", pageCompression=1,
                   invariant=1)

    def decoration(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(colors.black)
        canvas.setFont("CrimsonText-Bold", 8)
        canvas.drawCentredString(306, 755, BANNER)
        canvas.line(47, 748, 565, 748)
        canvas.setFont("CrimsonText", 9)
        canvas.drawCentredString(306, 32, str(document.page))
        canvas.restoreState()

    def line_numbers(canvas, document):
        if not b2 or document.page <= 1:
            return
        canvas.saveState()
        canvas.setFont("CrimsonText", 6.5)
        for column, x in (("left", 47), ("right", 318)):
            lowest_text = getattr(document, "_column_extents", {}).get(column)
            if lowest_text is None:
                continue
            for number in range(5, 56, 5):
                y = 706 - (number - 1) * 11.6
                if y > 57 and y >= lowest_text - 5:
                    canvas.drawRightString(x, y, str(number))
        canvas.restoreState()
        document._column_extents = {}

    if b2:
        first = Frame(47, 58, 518, 675, leftPadding=3, rightPadding=3,
                      topPadding=10, bottomPadding=4, id="cover")
        # 17pt structural gutter with reserved numeral space within the
        # inter-column whitespace; both columns have comparable text width.
        left = Frame(50, 60, 255, 666, leftPadding=4, rightPadding=5,
                     topPadding=10, bottomPadding=4, id="left")
        right = Frame(322, 60, 243, 666, leftPadding=4, rightPadding=5,
                      topPadding=10, bottomPadding=4, id="right")
        doc.addPageTemplates([PageTemplate(id="cover", frames=[first], onPage=decoration),
                              PageTemplate(id="body", frames=[left, right], onPage=decoration,
                                           onPageEnd=line_numbers)])
    else:
        frame = Frame(72, 59, 486, 655, leftPadding=0, rightPadding=0,
                      topPadding=2, bottomPadding=0, id="application")
        doc.addPageTemplates(PageTemplate(id="application", frames=[frame], onPage=decoration))

    def paragraph(item: Part):
        style = subheading if item.kind == "heading" else (status if item.kind == "status"
                  else claim if item.kind == "claim" else body)
        return Paragraph(escape(item.text).replace("\n", "<br/>"), style)

    story = []
    if b2:
        cover = Table([
            [Paragraph("(54) TITLE OF THE INVENTION", cover_label),
             Paragraph("(57) ABSTRACT", cover_label)],
            [Paragraph(escape(model.title), cover_text),
             Paragraph(escape(model.abstract), cover_text)],
        ], colWidths=[255, 255], hAlign="CENTER")
        cover.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("LINEABOVE", (0, 0), (-1, 0), 1, colors.black),
                                   ("LINEBELOW", (0, 1), (-1, 1), 0.6, colors.black),
                                   ("LEFTPADDING", (0, 0), (-1, -1), 9),
                                   ("RIGHTPADDING", (0, 0), (-1, -1), 9)]))
        story.extend([
            Paragraph(escape("DRAFT ID: " + model.draft_id),
                      ParagraphStyle("DraftId", parent=cover_label, fontSize=8,
                                     leading=10, spaceBefore=0, spaceAfter=20)),
            Spacer(1, 16), cover, NextPageTemplate("body"), PageBreak(),
        ])
    for section in model.sections:
        if not b2 and section.section_id in ("CLAIMS", "ABSTRACT_OF_THE_DISCLOSURE"):
            story.append(PageBreak())
        story.append(Paragraph(escape(section.heading), heading))
        for part in section.parts:
            story.append(paragraph(part))
    doc.build(story)


def _docx(model: RenderModel, path: Path, *, b2: bool) -> None:
    from docx import Document
    from docx.enum.section import WD_SECTION_START
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Crimson Text"
    normal.font.size = Pt(9.4 if b2 else 12)
    normal.paragraph_format.line_spacing = 1.12 if b2 else 1.5
    normal.paragraph_format.space_after = Pt(4 if b2 else 8)

    def setup(section, *, columns: int):
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(.9 if b2 else 1.0)
        section.bottom_margin = Inches(.7 if b2 else .75)
        section.left_margin = Inches(.66 if b2 else 1.0)
        section.right_margin = Inches(.66 if b2 else .75)
        section.header_distance = Inches(.32)
        section.footer_distance = Inches(.35)
        cols = section._sectPr.xpath("./w:cols")
        element = cols[0] if cols else OxmlElement("w:cols")
        element.set(qn("w:num"), str(columns))
        element.set(qn("w:space"), "840" if columns == 2 else "0")
        if not cols:
            section._sectPr.append(element)
        if b2 and columns == 2:
            lines = OxmlElement("w:lnNumType")
            for key, value in (("countBy", "5"), ("start", "1"),
                               ("distance", "160"), ("restart", "newPage")):
                lines.set(qn("w:" + key), value)
            section._sectPr.append(lines)

    setup(document.sections[0], columns=1)
    header = document.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.add_run(BANNER).bold = True
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def add_heading(text: str, *, small: bool = False):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(11 if not small else 7)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(10 if b2 else 12)

    def add_part(part: Part):
        if part.kind == "heading":
            add_heading(part.text, small=True)
            return
        p = document.add_paragraph()
        p.paragraph_format.widow_control = True
        if part.kind in ("claim", "status"):
            p.paragraph_format.keep_together = True
        if part.kind == "status":
            p.add_run(part.text).bold = True
        else:
            p.add_run(part.text)
        if b2 and part.kind != "status":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    if b2:
        document.add_paragraph(BANNER).runs[0].bold = True
        document.add_paragraph("DRAFT ID: " + model.draft_id)
        table = document.add_table(rows=2, cols=2)
        for row in table.rows:
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        table.cell(0, 0).text = "(54) TITLE OF THE INVENTION"
        table.cell(0, 1).text = "(57) ABSTRACT"
        table.cell(1, 0).text = model.title
        table.cell(1, 1).text = model.abstract
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
        body_section = document.add_section(WD_SECTION_START.NEW_PAGE)
        setup(body_section, columns=2)
        body_section.header.is_linked_to_previous = True
        body_section.footer.is_linked_to_previous = True

    for section in model.sections:
        if not b2 and section.section_id in ("CLAIMS", "ABSTRACT_OF_THE_DISCLOSURE"):
            document.add_page_break()
        add_heading(section.heading)
        for part in section.parts:
            add_part(part)
    document.core_properties.title = BANNER
    document.core_properties.subject = "Synthetic draft only; human IP review required"
    document.save(path)
